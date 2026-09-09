"""
syno_enrich.py – Syno-Importdatei vor dem Import SICHER mit DB-Daten anreichern.

Leitprinzip: Es wird nie eine unsichere Korrektur still in die Datei geschrieben.
Nur eindeutige, hochsichere Treffer werden direkt eingetragen; alles Unsichere
bleibt im Original stehen und wird als *Vorschlag* daneben markiert. Keine Zeile
wird entfernt – die Datei behält exakt ihre Zeilen, nur zusätzliche Hinweis-
Spalten und ein Bericht-Blatt kommen hinzu.

Vertrauensstufen pro Zeile:
  SICHER (grün/blau) – wird direkt eingetragen:
    - GSM in der Datei exakt in der DB gefunden  → Name aus DB korrigieren
    - Name exakt und EINDEUTIG in der DB gefunden → GSM aus DB eintragen
  VORSCHLAG (gelb) – wird NICHT eingetragen, nur vorgeschlagen:
    - nur ein unscharfer Namenstreffer (Zusatz, Teilname)
    - mehrere Namenskandidaten (mehrdeutig)
    - kein Treffer
Der anschließende Import verarbeitet die Datei vollständig; unsichere Zeilen
gehen dort nicht verloren, sondern landen in „Nicht zugeordnet".
"""

import logging
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill, Font

from modules import database as db
from modules import settings_store
from modules.utils import normalize_gsm, normalize_name, names_match, name_subset_match

logger = logging.getLogger(__name__)

FILL_FIXED   = PatternFill("solid", fgColor="C6EFCE")  # grün  – sicher ergänzt
FILL_CHANGED = PatternFill("solid", fgColor="BDD7EE")  # blau  – sicher korrigiert
FILL_WARN    = PatternFill("solid", fgColor="FFEB9C")  # gelb  – unsicher, bitte prüfen
FILL_PROPOSE = PatternFill("solid", fgColor="FCE4D6")  # orange – Vorschlag (nicht angew(andt)


class EnrichInputError(Exception):
    """Eingabedatei ist ungeeignet – es wird bewusst nichts verändert."""


def _raw_str(value) -> str:
    """Zellwert als String, ohne Float-Artefakte (Excel legt Nummern als Zahl ab).

    str(1735123456.0) -> '1735123456.0' – das angehängte '.0' erzeugte sonst eine
    falsche Extra-Ziffer bei der GSM-Normalisierung.
    """
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def run_syno_enrich(filepath: str | Path) -> dict:
    filepath = Path(filepath)
    if filepath.suffix.lower() != ".xlsx":
        raise EnrichInputError(
            "Nur .xlsx wird unterstützt (openpyxl liest kein .xls). "
            "Bitte die Datei in Excel als .xlsx speichern.")
    if not filepath.exists():
        raise FileNotFoundError(f"Datei nicht gefunden: {filepath}")

    cols = settings_store.get()["syno_columns"]
    col_gsm  = cols["rufnummer"]
    col_name = cols["name"]

    try:
        wb = openpyxl.load_workbook(str(filepath))
    except Exception as exc:
        raise EnrichInputError(f"Datei konnte nicht gelesen werden: {exc}") from exc
    ws = wb.active

    # --- Validierung, bevor irgendetwas angefasst wird ---------------------
    if ws.max_row < 2:
        raise EnrichInputError("Die Datei enthält keine Datenzeilen.")
    needed = max(col_gsm, col_name)
    if ws.max_column <= needed:
        raise EnrichInputError(
            f"Die konfigurierte Spaltenzuordnung passt nicht zur Datei "
            f"(erwartet mind. {needed + 1} Spalten, gefunden {ws.max_column}). "
            f"Bitte die Syno-Spalten in den Einstellungen prüfen.")
    # Plausibilität: sieht die GSM-Spalte über die Datei hinweg nach Nummern aus?
    gsm_seen = gsm_numeric = 0
    for row in ws.iter_rows(min_row=2, max_row=min(ws.max_row, 200)):
        v = _raw_str(row[col_gsm].value) if col_gsm < len(row) else ""
        if v:
            gsm_seen += 1
            if normalize_gsm(v):
                gsm_numeric += 1
    if gsm_seen and gsm_numeric / gsm_seen < 0.5:
        raise EnrichInputError(
            "Die als GSM konfigurierte Spalte enthält überwiegend keine "
            "Rufnummern. Vermutlich ist die Spaltenzuordnung verschoben – "
            "es wurde nichts verändert. Bitte die Syno-Spalten prüfen.")

    with db.read_connection() as conn:
        all_participants = db.get_all_participants(conn, provider=None)

    gsm_index = {}
    for p in all_participants:
        if p["gsm"]:
            ng = normalize_gsm(p["gsm"])
            if ng:
                gsm_index[ng] = p

    # Zusatzspalten hinter den Daten (verschieben die Import-Spalten nicht).
    col_vorschlag = ws.max_column + 1
    col_hinweis   = ws.max_column + 2
    ws.cell(1, col_vorschlag, "Vorschlag (nicht übernommen)").font = Font(bold=True)
    ws.cell(1, col_hinweis, "Hinweis").font = Font(bold=True)

    report = []   # (zeile, aktion, alt, neu_oder_vorschlag, vertrauen)
    c = {"ok_gsm": 0, "fixed_gsm": 0, "changed_gsm": 0, "fixed_name": 0,
         "vorschlag": 0, "mehrdeutig": 0, "no_match": 0, "rows_total": 0}

    for row in ws.iter_rows(min_row=2):
        rownum = row[0].row
        raw_gsm  = _raw_str(row[col_gsm].value)  if col_gsm  < len(row) else ""
        raw_name = _raw_str(row[col_name].value) if col_name < len(row) else ""
        if not raw_gsm and not raw_name:
            continue  # echte Leerzeile – nichts zu tun, wird nicht gezählt
        c["rows_total"] += 1

        gsm_norm  = normalize_gsm(raw_gsm)
        name_norm = normalize_name(raw_name)
        hinweis_cell = row[col_hinweis - 1] if col_hinweis - 1 < len(row) else ws.cell(rownum, col_hinweis)

        # 1) SICHER: GSM exakt in der DB → Name korrigieren/ergänzen (grün)
        if gsm_norm and gsm_norm in gsm_index:
            matched = gsm_index[gsm_norm]
            db_name = (matched["name"] or "").strip()
            if db_name and raw_name.strip() != db_name:
                row[col_name].value = db_name
                row[col_name].fill  = FILL_FIXED
                c["fixed_name"] += 1
                report.append((rownum, "Name korrigiert (GSM-Treffer)", raw_name, db_name, "sicher"))
            else:
                c["ok_gsm"] += 1
            ws.cell(rownum, col_hinweis, "GSM in Bestand gefunden").fill = FILL_FIXED
            continue

        # 2) Namensabgleich
        if name_norm:
            exact = [p for p in all_participants
                     if names_match(name_norm, normalize_name(p["name"] or ""))]
            if len(exact) == 1:
                # SICHER: exakter, eindeutiger Name → GSM eintragen (grün/blau)
                matched = exact[0]
                db_gsm = normalize_gsm(matched["gsm"] or "")
                if db_gsm:
                    if raw_gsm.strip() and normalize_gsm(raw_gsm) != db_gsm:
                        row[col_gsm].value = db_gsm
                        row[col_gsm].fill  = FILL_CHANGED
                        c["changed_gsm"] += 1
                        report.append((rownum, "GSM korrigiert (Name eindeutig)", raw_gsm, db_gsm, "sicher"))
                    else:
                        row[col_gsm].value = db_gsm
                        row[col_gsm].fill  = FILL_FIXED
                        c["fixed_gsm"] += 1
                        report.append((rownum, "GSM ergänzt (Name eindeutig)", raw_gsm or "—", db_gsm, "sicher"))
                    ws.cell(rownum, col_hinweis, "Name eindeutig im Bestand").fill = FILL_FIXED
                    continue
            elif len(exact) > 1:
                # MEHRDEUTIG: nichts eintragen, nur markieren
                _flag(row, col_gsm, col_name, FILL_WARN)
                ws.cell(rownum, col_hinweis, f"Mehrdeutig: {len(exact)} Namenskandidaten – bitte prüfen").fill = FILL_WARN
                c["mehrdeutig"] += 1
                report.append((rownum, "Mehrdeutig", raw_name, f"{len(exact)} Kandidaten", "unsicher"))
                continue

            # VORSCHLAG: nur unscharfer, eindeutiger Namenstreffer → NICHT eintragen
            fuzzy = [p for p in all_participants
                     if name_subset_match(name_norm, normalize_name(p["name"] or ""))]
            if len(fuzzy) == 1:
                cand = fuzzy[0]
                db_gsm = normalize_gsm(cand["gsm"] or "")
                vorschlag = f"GSM {db_gsm or '—'} (ähnlich: '{(cand['name'] or '').strip()}')"
                ws.cell(rownum, col_vorschlag, vorschlag).fill = FILL_PROPOSE
                _flag(row, col_gsm, col_name, FILL_WARN)
                ws.cell(rownum, col_hinweis, "Unsicher: unscharfer Namenstreffer – Vorschlag prüfen").fill = FILL_WARN
                c["vorschlag"] += 1
                report.append((rownum, "Vorschlag (unscharf, NICHT übernommen)", raw_name, vorschlag, "unsicher"))
                continue

        # 3) KEIN TREFFER: markieren, Originaldaten unangetastet lassen
        _flag(row, col_gsm, col_name, FILL_WARN)
        ws.cell(rownum, col_hinweis, "Kein Treffer im Bestand – bitte prüfen").fill = FILL_WARN
        c["no_match"] += 1
        report.append((rownum, "Kein Treffer", raw_gsm or raw_name, "—", "unsicher"))

    _write_legend(ws, col_hinweis)
    _write_report(wb, report, c)

    out_path = filepath.with_stem(filepath.stem + "_angereichert")
    wb.save(str(out_path))

    summary = dict(c)
    summary["out_path"] = out_path
    # Rückwärtskompatible Schlüssel für den API-Handler:
    summary.setdefault("fixed_gsm", 0)
    summary.setdefault("changed_gsm", 0)
    summary.setdefault("fixed_name", 0)
    summary.setdefault("no_match", 0)
    summary["unsicher"] = c["vorschlag"] + c["mehrdeutig"] + c["no_match"]
    logger.info(
        "Syno-Anreicherung (sicher): %d Zeilen, sicher: %d GSM ergänzt / %d GSM "
        "korrigiert / %d Namen; unsicher: %d Vorschlag / %d mehrdeutig / %d ohne "
        "Treffer → %s",
        c["rows_total"], c["fixed_gsm"], c["changed_gsm"], c["fixed_name"],
        c["vorschlag"], c["mehrdeutig"], c["no_match"], out_path.name,
    )
    return summary


def _flag(row, col_gsm, col_name, fill):
    for idx in (col_gsm, col_name):
        if idx < len(row):
            row[idx].fill = fill


def _write_legend(ws, base_col):
    col = base_col + 2
    ws.cell(1, col, "Legende:").font = Font(bold=True)
    ws.cell(2, col, "Sicher ergänzt").fill      = FILL_FIXED
    ws.cell(3, col, "Sicher korrigiert").fill   = FILL_CHANGED
    ws.cell(4, col, "Unsicher – bitte prüfen").fill = FILL_WARN
    ws.cell(5, col, "Vorschlag (nicht übernommen)").fill = FILL_PROPOSE


def _write_report(wb, report, counts):
    sh = wb.create_sheet("Bericht")
    sh.append(["Zeile", "Aktion", "Alt", "Neu / Vorschlag", "Vertrauen"])
    for cell in sh[1]:
        cell.font = Font(bold=True)
    for zeile, aktion, alt, neu, vertrauen in report:
        sh.append([zeile, aktion, alt, neu, vertrauen])
    sh.append([])
    sh.append(["Zusammenfassung", "", "", "", ""])
    sh.append(["Zeilen gesamt", counts["rows_total"]])
    sh.append(["Sicher: GSM ergänzt", counts["fixed_gsm"]])
    sh.append(["Sicher: GSM korrigiert", counts["changed_gsm"]])
    sh.append(["Sicher: Name korrigiert", counts["fixed_name"]])
    sh.append(["Unsicher: Vorschlag (unscharf)", counts["vorschlag"]])
    sh.append(["Unsicher: mehrdeutig", counts["mehrdeutig"]])
    sh.append(["Unsicher: kein Treffer", counts["no_match"]])
