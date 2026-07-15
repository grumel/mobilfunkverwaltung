"""
syno_import.py – Monatlicher Geräteimport aus Syno-Exportdatei.

Importiert nur: Gerätemodell und Startdatum, in bis zu zwei Slots je
Teilnehmer (syno/start_syno und syno2/start_syno2). Eine Person mit zwei
Geräten steht in zwei Zeilen (gleiche Nummer) – die Geräte füllen nach-
einander die freien Slots. Bereits vorhandene Geräte werden nicht doppelt
eingetragen, belegte Slots nicht überschrieben ("nur leere Slots füllen").

Matching-Reihenfolge:
  1. Rufnummer/GSM (Spalte P, Index 15)
  2. Name (Spalte Q, Index 16) – nur bei eindeutigem Treffer

Kein Treffer:
  -> Eintrag in unmatched_devices (Tab "Nicht zugeordnet")
     Dort kann das Gerät manuell einem Teilnehmer zugewiesen werden.
"""

import logging
from pathlib import Path
from datetime import datetime

import openpyxl

from modules import database as db
from modules import settings_store
from modules.utils import normalize_gsm, normalize_date, normalize_name, names_match, name_subset_match

logger = logging.getLogger(__name__)


def _cell_value(row, col_idx) -> str | None:
    if col_idx >= len(row):
        return None
    v = row[col_idx].value
    if v is None:
        return None
    return str(v).strip() or None


def find_by_gsm_or_name(filepath: str | Path, gsm: str | None, name: str | None) -> dict | None:
    """Sucht die Geräte einer Person per GSM (bevorzugt) oder Name in der Syno-Datei.

    Da eine Person mit zwei Geräten in zwei Zeilen (gleiche Nummer) auftaucht,
    werden bis zu zwei Geräte gesammelt und als syno/start_syno bzw.
    syno2/start_syno2 zurückgegeben. Liest nur die Datei, schreibt nichts in die
    Datenbank – für den Einzel-Abgleich im Bearbeiten-Fenster. Gibt None zurück,
    wenn kein Treffer gefunden wurde.
    """
    filepath = Path(filepath)
    if not filepath.exists():
        raise FileNotFoundError(f"Syno-Datei nicht gefunden: {filepath}")

    cols = settings_store.get()["syno_columns"]
    wb = openpyxl.load_workbook(str(filepath), data_only=True)
    ws = wb.active

    target_gsm  = normalize_gsm(gsm) if gsm else None
    target_name = normalize_name(name) if name else None

    gsm_devices:  list[tuple[str, str | None]] = []
    name_devices: list[tuple[str, str | None]] = []

    for row in ws.iter_rows(min_row=2):
        raw_gsm  = _cell_value(row, cols["rufnummer"])
        raw_name = _cell_value(row, cols["name"])
        raw_syno = _cell_value(row, cols["syno"])
        raw_date = _cell_value(row, cols["startdatum"])
        if not raw_syno:
            continue

        if target_gsm and normalize_gsm(raw_gsm) == target_gsm:
            gsm_devices.append((raw_syno, normalize_date(raw_date)))
        elif target_name:
            row_name = normalize_name(raw_name)
            if names_match(target_name, row_name) or name_subset_match(target_name, row_name):
                name_devices.append((raw_syno, normalize_date(raw_date)))

    devices = gsm_devices or name_devices
    if not devices:
        return None
    # Duplikate (gleiches Gerät) entfernen, max. 2 behalten
    seen: set[str] = set()
    unique: list[tuple[str, str | None]] = []
    for dev, start in devices:
        key = dev.strip().casefold()
        if key not in seen:
            seen.add(key)
            unique.append((dev, start))
    result = {"syno": unique[0][0], "start_syno": unique[0][1]}
    if len(unique) > 1:
        result["syno2"] = unique[1][0]
        result["start_syno2"] = unique[1][1]
    return result


# ---------------------------------------------------------------------------
# Öffentliche Importfunktion
# ---------------------------------------------------------------------------

def run_syno_import(filepath: str | Path, db_module=None, create_missing=False) -> dict:
    """db_module: austauschbares Datenbank-Backend (Standard: modules.database,
    SQLite). Siehe webapp/import_adapter.py für die SQLAlchemy-Variante.

    create_missing: Zeilen ohne Treffer, die aber einen Namen oder eine GSM
    haben, werden als neue Teilnehmer angelegt (statt in 'Nicht zugeordnet')."""
    dbm = db_module or db
    filepath = Path(filepath)
    if not filepath.exists():
        raise FileNotFoundError(f"Syno-Datei nicht gefunden: {filepath}")

    logger.info("Syno-Import gestartet: %s", filepath)

    wb = openpyxl.load_workbook(str(filepath), data_only=True)
    ws = wb.active

    cols = settings_store.get()["syno_columns"]

    matched_gsm   = 0
    matched_name  = 0
    unmatched     = 0
    skipped       = 0
    duplicate     = 0   # Gerät war bereits eingetragen
    slots_full    = 0   # beide Geräte-Slots belegt
    created_new   = 0   # neu angelegte Teilnehmer (create_missing)
    errors: list[str] = []
    log_lines: list[str] = []

    log_lines.append(f"=== Syno-Import {datetime.now().strftime('%d.%m.%Y %H:%M:%S')} ===")
    log_lines.append(f"Datei: {filepath.name}")

    with dbm.transaction() as conn:
        # provider=None: auch Telekom-Teilnehmer für das Matching heranziehen
        all_participants = dbm.get_all_participants(conn, provider=None)
        # Normalisierter GSM-Index: robustes Matching unabhängig von der
        # Schreibweise (Leerzeichen/Formate) – verhindert doppelte Neuanlagen.
        gsm_index = {}
        for _p in all_participants:
            _ng = normalize_gsm(_p["gsm"] or "")
            if _ng:
                gsm_index.setdefault(_ng, _p)

        for row in ws.iter_rows(min_row=2):
            raw_gsm  = _cell_value(row, cols["rufnummer"])
            raw_name = _cell_value(row, cols["name"])
            raw_syno = _cell_value(row, cols["syno"])
            raw_date = _cell_value(row, cols["startdatum"])

            gsm        = normalize_gsm(raw_gsm)
            syno       = raw_syno
            start_syno = normalize_date(raw_date)

            if not syno and not gsm and not raw_name:
                skipped += 1
                continue

            def _apply_device(pid: int, quelle_aktion: str, kontext: str) -> str:
                """Gerät in freien Slot eintragen und Ergebnis protokollieren."""
                res = dbm.assign_syno_device(conn, pid, syno, start_syno)
                if res == "filled":
                    dbm.log_import(conn, "Syno", quelle_aktion,
                                  f"ID={pid} {kontext} Syno={syno}")
                elif res == "full":
                    log_lines.append(
                        f"SLOTS BELEGT: ID={pid} {kontext} – Gerät '{syno}' "
                        f"nicht eingetragen (beide Geräte-Slots belegt)"
                    )
                return res

            try:
                # 1. Matching über GSM (normalisiert, robust gegen Formate)
                if gsm:
                    match = gsm_index.get(gsm)
                    if match:
                        res = _apply_device(match["id"], "UPDATE_GSM", f"GSM={gsm}")
                        if res == "filled":
                            matched_gsm += 1
                        elif res == "dup":
                            duplicate += 1
                        else:
                            slots_full += 1
                        continue

                # 2. Matching über normalisierten Namen (nur eindeutig)
                norm_import = normalize_name(raw_name)
                if norm_import:
                    candidates = [
                        p for p in all_participants
                        if names_match(norm_import, normalize_name(p["name"]))
                    ]
                    if not candidates:
                        # Toleranter Rückfall: Name mit Zusatz ("Kessel Ralf ehem. …")
                        candidates = [
                            p for p in all_participants
                            if name_subset_match(norm_import, normalize_name(p["name"]))
                        ]
                    if len(candidates) == 1:
                        res = _apply_device(candidates[0]["id"], "UPDATE_NAME",
                                            f"Name={raw_name}")
                        if res == "filled":
                            log_lines.append(
                                f"Name-Match: '{raw_name}' → ID={candidates[0]['id']}"
                            )
                            matched_name += 1
                        elif res == "dup":
                            duplicate += 1
                        else:
                            slots_full += 1
                        continue
                    elif len(candidates) > 1:
                        log_lines.append(
                            f"MEHRDEUTIG: '{raw_name}' – {len(candidates)} Treffer, "
                            f"in 'Nicht zugeordnet' abgelegt"
                        )

                # 3. Kein Treffer → neuen Teilnehmer anlegen (falls gewünscht)
                #    oder in unmatched_devices ablegen.
                if create_missing and ((raw_name or "").strip() or gsm):
                    new_id = dbm.insert_participant(conn, {
                        "name":           (raw_name or "").strip() or None,
                        "gsm":            gsm or None,
                        "syno":           syno or None,
                        "start_syno":     start_syno or None,
                        "provider":       "Vodafone",
                        "verified":       0,
                        "pruefung_grund": "Neu aus Syno-Import angelegt",
                        "bemerkung":      "Automatisch aus Syno-Import (kein Treffer im Bestand)",
                    })
                    dbm.log_import(conn, "Syno", "NEU_ANGELEGT",
                        f"ID={new_id} Name={raw_name} GSM={gsm} Syno={syno} – neu aus Syno")
                    log_lines.append(
                        f"NEU ANGELEGT: Name={raw_name} GSM={gsm} Syno={syno} → ID={new_id} (zur Prüfung)"
                    )
                    # In die Indizes aufnehmen, damit weitere Zeilen derselben
                    # Person nicht erneut angelegt werden.
                    new_p = {"id": new_id, "name": (raw_name or "").strip(),
                             "gsm": gsm, "syno": syno, "syno2": None}
                    all_participants.append(new_p)
                    if gsm:
                        gsm_index.setdefault(gsm, new_p)
                    created_new += 1
                    continue

                dbm.insert_unmatched_device(conn, {
                    "quelle":    "Syno",
                    "gsm":       gsm,
                    "benutzer":  raw_name,
                    "geraet":    syno,
                    "startdatum": start_syno,
                })
                dbm.log_import(conn, "Syno", "UNMATCHED",
                    f"GSM={gsm} Name={raw_name} Syno={syno} – kein Treffer")
                log_lines.append(
                    f"KEIN TREFFER: GSM={gsm} Name={raw_name} Syno={syno} → 'Nicht zugeordnet'"
                )
                unmatched += 1

            except Exception as exc:
                # Unsichere Zeile NICHT verlieren: als Fehler zur Prüfung ablegen.
                row_num = row[0].row if row else "?"
                msg = f"Zeile {row_num}: {exc}"
                logger.error(msg)
                errors.append(msg)
                try:
                    dbm.insert_unmatched_device(conn, {
                        "quelle":    "Syno-Fehler",
                        "gsm":       gsm,
                        "benutzer":  raw_name,
                        "geraet":    f"Importfehler: {exc}",
                        "startdatum": None,
                    })
                except Exception as exc2:
                    logger.error("Zeile %s konnte nicht zur Prüfung abgelegt "
                                 "werden: %s", row_num, exc2)
                skipped += 1

        summary = (
            f"Syno-Import abgeschlossen: {matched_gsm} GSM-Matches, "
            f"{matched_name} Name-Matches, {created_new} neu angelegt, "
            f"{duplicate} bereits vorhanden, {slots_full} ohne freien Slot, "
            f"{unmatched} nicht zugeordnet, {skipped} übersprungen, {len(errors)} Fehler"
        )
        dbm.log_import(conn, "Syno", "SUMMARY", summary)
        log_lines.append(summary)

    if errors:
        log_lines.append("--- Fehler (zur Prüfung in 'Nicht zugeordnet') ---")
        log_lines.extend(errors)

    logger.info(summary)
    return {
        "matched_gsm":  matched_gsm,
        "matched_name": matched_name,
        "neu_angelegt": created_new,
        "duplicate":    duplicate,
        "slots_full":   slots_full,
        "unmatched":    unmatched,
        "skipped":      skipped,
        "errors":       errors,
        "log_lines":    log_lines,
    }
