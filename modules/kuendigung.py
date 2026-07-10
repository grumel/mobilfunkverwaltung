"""
kuendigung.py – Kündigungs-/Rücknahme-Schreiben aus Word-Vorlage erzeugen,
als PDF speichern und als Outlook-Entwurf mit Anhang öffnen.

Voraussetzungen auf dem ausführenden Rechner:
  - Microsoft Word (für die PDF-Konvertierung über COM)
  - Microsoft Outlook Desktop (für den E-Mail-Entwurf über COM)

Platzhalter in den Vorlagen: {{nummer}} – wird durch die GSM-Nummer ersetzt.
Das Datum ("Berlin, DD.MM.YYYY") wird automatisch auf das heutige Datum
aktualisiert.
"""

import logging
import re
from pathlib import Path
from datetime import datetime, date

import docx

from modules.paths import DATA_DIR

logger = logging.getLogger(__name__)

BASE_DIR      = DATA_DIR
TEMPLATE_DIR  = BASE_DIR / "Dokumente"
OUTPUT_DIR    = BASE_DIR / "Kuendigungen"

TEMPLATES = {
    "kuendigung": TEMPLATE_DIR / "vorlage_kündigung.docx",
    "ruecknahme": TEMPLATE_DIR / "vorlage_rücknahme.docx",
}

SUBJECTS = {
    "kuendigung": "Kündigung GSM {gsm}",
    "ruecknahme": "Rücknahme Kündigung GSM {gsm}",
}

DEFAULT_TO = "nicole.dieckmann@zinxs.com"

BODIES = {
    "kuendigung": (
        "Hallo liebe Nicole,\n\n"
        "anbei die Kündigung für die GSM-Nummer {gsm} – magst du dich bitte "
        "wieder darum kümmern?\n\n"
        "Danke dir und ganz liebe Grüße\n"
        "Holger"
    ),
    "ruecknahme": (
        "Hallo liebe Nicole,\n\n"
        "anbei die Rücknahme der Kündigung für die GSM-Nummer {gsm} – magst "
        "du dich bitte wieder darum kümmern?\n\n"
        "Danke dir und ganz liebe Grüße\n"
        "Holger"
    ),
}

_DATE_RE = re.compile(r"\d{2}\.\d{2}\.\d{4}")


def _replace_placeholder_in_paragraph(paragraph, placeholder: str, value: str) -> bool:
    """Ersetzt einen Platzhalter, auch wenn er über mehrere Word-Runs verteilt ist."""
    if placeholder not in paragraph.text or not paragraph.runs:
        return False
    new_text = paragraph.text.replace(placeholder, value)
    paragraph.runs[0].text = new_text
    for r in paragraph.runs[1:]:
        r.text = ""
    return True


def _update_date_paragraph(paragraph) -> None:
    """Ersetzt ein Datum im Format DD.MM.YYYY durch das heutige Datum."""
    if not paragraph.runs or not _DATE_RE.search(paragraph.text):
        return
    new_text = _DATE_RE.sub(date.today().strftime("%d.%m.%Y"), paragraph.text)
    if new_text != paragraph.text:
        paragraph.runs[0].text = new_text
        for r in paragraph.runs[1:]:
            r.text = ""


def generate_letter(kind: str, gsm: str) -> Path:
    """Füllt die Vorlage aus und speichert sie als neue .docx. Gibt den Pfad zurück."""
    if kind not in TEMPLATES:
        raise ValueError(f"Unbekannter Schreiben-Typ: {kind}")
    template = TEMPLATES[kind]
    if not template.exists():
        raise FileNotFoundError(f"Vorlage nicht gefunden: {template}")

    gsm = (gsm or "").strip() or "unbekannt"
    doc = docx.Document(str(template))
    for p in doc.paragraphs:
        _replace_placeholder_in_paragraph(p, "{{nummer}}", gsm)
        _update_date_paragraph(p)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _replace_placeholder_in_paragraph(p, "{{nummer}}", gsm)

    OUTPUT_DIR.mkdir(exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_path = OUTPUT_DIR / f"{kind}_{gsm}_{ts}.docx"
    doc.save(str(out_path))
    logger.info("Schreiben erzeugt: %s", out_path.name)
    return out_path


def convert_to_pdf(docx_path: Path) -> Path:
    """Konvertiert eine .docx über installiertes MS Word (COM) zu PDF.

    Nutzt Dispatch() statt DispatchEx(): eine isolierte Word-Instanz
    (DispatchEx) kann bei OneDrive-Pfaden mit Leerzeichen/Sonderzeichen
    keinen Zugriff auf die Datei bekommen ("Datei nicht gefunden").
    """
    import time
    import win32com.client

    docx_path = docx_path.resolve()
    pdf_path = docx_path.with_suffix(".pdf")
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        wdoc = word.Documents.Open(
            str(docx_path), ReadOnly=False, ConfirmConversions=False,
            AddToRecentFiles=False,
        )
        try:
            # Vereinzelt ist das COM-Objekt direkt nach Open() noch nicht
            # vollständig bereit (SaveAs wirft dann AttributeError) – kurz
            # warten und erneut versuchen.
            last_exc = None
            for attempt in range(3):
                try:
                    wdoc.SaveAs(str(pdf_path), FileFormat=17)  # wdFormatPDF
                    last_exc = None
                    break
                except AttributeError as exc:
                    last_exc = exc
                    logger.warning("SaveAs-Versuch %d fehlgeschlagen, erneut …", attempt + 1)
                    time.sleep(0.7)
            if last_exc:
                raise last_exc
        finally:
            wdoc.Close(False)
    finally:
        word.Quit()
    return pdf_path


def convert_to_pdf_soffice(docx_path: Path) -> Path:
    """Konvertiert eine .docx über LibreOffice (headless) zu PDF – für den
    Linux-Server (statt Word-COM). Nutzt ein eigenes, temporäres LibreOffice-
    Profil, damit es nicht mit einer offenen LibreOffice-Sitzung kollidiert."""
    import shutil
    import subprocess
    import tempfile

    docx_path = Path(docx_path).resolve()
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError("LibreOffice (soffice) nicht gefunden – bitte installieren.")

    outdir = docx_path.parent
    with tempfile.TemporaryDirectory(prefix="mobilfunk_soffice_") as profile:
        cmd = [
            soffice,
            f"-env:UserInstallation=file://{profile}",
            "--headless", "--nologo", "--nofirststartwizard",
            "--convert-to", "pdf", "--outdir", str(outdir), str(docx_path),
        ]
        proc = subprocess.run(cmd, capture_output=True, text=True, timeout=120)
    pdf_path = docx_path.with_suffix(".pdf")
    if not pdf_path.exists():
        raise RuntimeError(
            "PDF-Erzeugung fehlgeschlagen"
            + (f": {proc.stderr.strip()}" if proc.stderr else ".")
        )
    logger.info("PDF erzeugt (LibreOffice): %s", pdf_path.name)
    return pdf_path


def create_outlook_draft(pdf_path: Path, subject: str, body: str = "", to: str = "") -> None:
    """Öffnet einen Outlook-Mail-Entwurf mit PDF-Anhang (nicht automatisch gesendet)."""
    import win32com.client

    outlook = win32com.client.Dispatch("Outlook.Application")
    mail = outlook.CreateItem(0)  # olMailItem
    mail.Subject = subject
    if body:
        mail.Body = body
    if to:
        mail.To = to
    mail.Attachments.Add(str(pdf_path))
    mail.Display()


def create_and_open(kind: str, gsm: str, name: str = "", to: str = "") -> Path:
    """Kompletter Ablauf: Vorlage füllen → PDF erzeugen → Outlook-Entwurf öffnen."""
    docx_path = generate_letter(kind, gsm)
    pdf_path = convert_to_pdf(docx_path)
    gsm_txt = gsm or "unbekannt"
    subject = SUBJECTS[kind].format(gsm=gsm_txt)
    body = BODIES[kind].format(gsm=gsm_txt)
    create_outlook_draft(pdf_path, subject=subject, body=body, to=to or DEFAULT_TO)
    return pdf_path
